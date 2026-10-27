#!/usr/bin/env python3
"""Text/table coverage log, not automatic proof of absence."""
from pathlib import Path
import re,csv,json,zipfile,io,hashlib
from pypdf import PdfReader
from openpyxl import load_workbook
P=Path(__file__).resolve().parents[2];O=P/'outputs/round2_A1';C=P/'work/round2_A1/novelty_receipts'
PAT=re.compile(r'\bintraclass\b|\bICC\b|\bMDC\b|(?:minimum|minimal) detectable|test[– -]retest|repeatability|regression to the mean',re.I)
def xlsx(p):
 b=load_workbook(p,read_only=True,data_only=True);return '\n'.join(' '.join(str(v) for v in row if v is not None) for s in b for row in s.iter_rows(values_only=True))
def main():
 files=list((P/'work/sources').glob('41586_2025_9686*'))+[P/'work/sources/gong.pdf',P/'work/extracted/gong.html.txt']+list(C.glob('IDEAL_Age*'));rows=[];hits=[]
 for p in files:
  row={'file':str(p.relative_to(P)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'status':'TEXT_SCANNED'}
  try:
   if p.suffix=='.pdf':
    pdf=PdfReader(p);row['pages']=len(pdf.pages);txt='\n'.join(page.extract_text() or '' for page in pdf.pages);(C/(p.stem+'_extracted.txt')).write_text(txt)
   elif p.suffix=='.xlsx':txt=xlsx(p)
   elif p.suffix=='.zip':
    with zipfile.ZipFile(p) as z:txt='\n'.join(xlsx(io.BytesIO(z.read(n))) for n in z.namelist() if n.endswith('.xlsx') and not n.startswith('__MACOSX'))
   elif p.suffix=='.txt':txt=p.read_text()
   else:continue
   ms=list(PAT.finditer(txt));row['n_keyword_hits']=len(ms);row['n_text_characters']=len(txt)
   for m in ms:hits.append({'file':row['file'],'term':m.group(),'context':txt[max(0,m.start()-100):m.end()+180].replace('\n',' ')})
  except Exception as e:row.update(status='UNREADABLE',error=str(e))
  rows.append(row)
 with (O/'novelty_source_coverage_QA.csv').open('w') as f:
  w=csv.DictWriter(f,fieldnames=sorted({k for r in rows for k in r}));w.writeheader();w.writerows(rows)
 with (O/'novelty_source_keyword_hits.csv').open('w') as f:
  w=csv.DictWriter(f,fieldnames=['file','term','context']);w.writeheader();w.writerows(hits)
 print(json.dumps(rows,indent=2));print('KEYWORD_HITS',len(hits))
if __name__=='__main__':main()
