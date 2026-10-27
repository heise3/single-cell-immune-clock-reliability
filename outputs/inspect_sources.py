#!/usr/bin/env python3
"""Extract public supplementary material and scan reliability terminology.

This only inventories publications/source files. It computes no clock outcomes.
Requires pypdf and openpyxl; seed reserved for any later K4 is 20261025.
"""
from pathlib import Path
from html.parser import HTMLParser
import json, re, zipfile, io
from pypdf import PdfReader
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
WORK=ROOT.parent/'work'/'sources'
TXT=ROOT.parent/'work'/'extracted'; TXT.mkdir(exist_ok=True)
PAT=re.compile(r'\bintraclass\b|\bICC\b|\bMDC\b|(?:minimum|minimal) detectable|test[– -]retest|repeatability|reliability',re.I)

class Plain(HTMLParser):
 def __init__(self): super().__init__(); self.parts=[]; self.skip=0
 def handle_starttag(self,t,a):
  if t in ['script','style']: self.skip+=1
  if t in ['p','h1','h2','h3','h4','li','figcaption','tr','div']: self.parts.append('\n')
 def handle_endtag(self,t):
  if t in ['script','style']: self.skip=max(0,self.skip-1)
 def handle_data(self,d):
  if not self.skip:self.parts.append(d)

def matches(s):
 return [{'term':m.group(),'context':s[max(0,m.start()-100):m.end()+180]} for m in PAT.finditer(s)]

def xlsx_info(data):
 book=load_workbook(data,read_only=True,data_only=True)
 out=[]
 for sheet in book:
  strings=[]; preview=[]
  for row in sheet.iter_rows(values_only=True):
   if len(preview)<3: preview.append([str(v)[:120] if v is not None else None for v in row[:10]])
   strings.extend(str(v) for v in row if isinstance(v,str))
  out.append({'sheet':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'preview':preview,'hits':matches('\n'.join(strings))})
 return out

out=[]
for p in sorted(WORK.iterdir()):
 item={'file':p.name}
 if p.suffix=='.pdf':
  try:
   pdf=PdfReader(p); pages=[page.extract_text() for page in pdf.pages]
   (TXT/(p.name+'.txt')).write_text('\n\n'.join(f'PAGE {i+1}\n{s}' for i,s in enumerate(pages)))
   item.update(pages=len(pages),hits=[{'page':i+1,**hit} for i,s in enumerate(pages) for hit in matches(s)])
  except Exception as e:item['error']=str(e)
 elif p.suffix=='.xlsx':item['sheets']=xlsx_info(p)
 elif p.suffix=='.zip':
  with zipfile.ZipFile(p) as archive:
   item['members']=[]
   for name in archive.namelist():
    if name.endswith('.xlsx') and not name.startswith('__MACOSX'):
     item['members'].append({'name':name,'sheets':xlsx_info(io.BytesIO(archive.read(name)))})
 elif p.suffix=='.html':
  parser=Plain();parser.feed(p.read_text()); s=''.join(parser.parts)
  (TXT/(p.name+'.txt')).write_text(s);item['hits']=matches(s)
 else:continue
 out.append(item)
(ROOT/'source_inspection.json').write_text(json.dumps(out,ensure_ascii=False,indent=2))
for item in out:
 print(item['file'],'pages',item.get('pages'),'hits',len(item.get('hits',[])))
 for hit in item.get('hits',[])[:8]: print(hit)
